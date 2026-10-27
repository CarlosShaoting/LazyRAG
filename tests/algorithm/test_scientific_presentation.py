from pptx import Presentation

from lazymind.chat.engine.tools.file_resources.scientific_presentation import build_scientific_pptx


def test_builds_editable_scientific_deck_with_sources_and_notes(tmp_path):
    output = tmp_path / 'science.pptx'
    result = build_scientific_pptx([
        {'title': 'Method X', 'bullets': ['Research briefing'], 'notes': 'Introduce the question.'},
        {'title': 'Method', 'bullets': ['Encode inputs', 'Compare baselines'],
         'equations': ['L = L_task + λL_reg'], 'sources': ['Paper A, §3'],
         'notes': 'Explain the method and its assumptions.'},
        {'title': 'Conclusion', 'bullets': ['Evidence supports the reported result'],
         'sources': ['Paper A, Table 2'], 'notes': 'Close with limitations.'},
    ], output)
    prs = Presentation(output)
    assert len(prs.slides) == 3
    assert prs.slides[1].shapes.title.text == 'Method'
    assert any('公式：L = L_task' in shape.text for shape in prs.slides[1].shapes if shape.has_text_frame)
    assert 'Paper A, §3' in prs.slides[1].notes_slide.notes_text_frame.text
    assert result['equations_as_text'] == 1
    assert any(item['code'] == 'EQUATIONS_AS_TEXT' for item in result['warnings'])


def test_reuses_template_without_sample_slides(tmp_path):
    template, output = tmp_path / 'template.pptx', tmp_path / 'result.pptx'
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = 'Sample content'
    prs.save(template)
    result = build_scientific_pptx([
        {'title': 'New cover', 'bullets': ['Grounded report']},
        {'title': 'Result', 'bullets': ['Finding'], 'sources': ['Paper B']},
    ], output, template=template)
    actual = Presentation(output)
    assert len(actual.slides) == 2
    assert actual.slides[0].shapes.title.text == 'New cover'
    assert 'Sample content' not in '\n'.join(
        shape.text for slide in actual.slides for shape in slide.shapes if shape.has_text_frame)
    assert result['template_mode'] == 'layout'


def test_rejects_empty_middle_slide(tmp_path):
    try:
        build_scientific_pptx([
            {'title': 'Cover', 'bullets': ['Intro']},
            {'title': 'Empty'},
            {'title': 'End', 'bullets': ['Done']},
        ], tmp_path / 'bad.pptx')
    except ValueError as exc:
        assert 'needs bullets' in str(exc)
    else:
        raise AssertionError('expected empty scientific content slide rejection')
