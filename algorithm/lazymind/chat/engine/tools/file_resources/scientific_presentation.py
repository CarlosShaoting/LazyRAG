"""Build editable scientific PPTX files for direct main-chat delivery."""
from __future__ import annotations

from pathlib import Path
from typing import Any


def _clean_text(value: Any, *, field: str, limit: int = 10000) -> str:
    text = str(value or '').strip()
    if not text:
        raise ValueError(f'{field} must be non-empty text')
    if len(text) > limit or any(ord(char) < 32 and char not in '\t\n\r' for char in text):
        raise ValueError(f'{field} contains invalid or excessive text')
    return text


def _clean_text_list(value: Any, *, field: str, limit: int = 100) -> list[str]:
    if value is None:
        return []
    if not isinstance(value, (list, tuple)) or len(value) > limit:
        raise ValueError(f'{field} must be a list with at most {limit} items')
    return [_clean_text(item, field=field) for item in value]


def build_scientific_pptx(
    slides: list[dict[str, Any]],
    output: Path,
    *,
    template: Path | None = None,
    image_paths: dict[str, Path] | None = None,
) -> dict[str, Any]:
    """Create an editable scientific deck without a Workflow.

    The main Chat is responsible for reading source material and grounding the
    supplied slide content. This writer only performs deterministic PPTX
    construction and keeps equations as editable text rather than pretending
    to produce native PowerPoint equations.
    """
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    if not isinstance(slides, list) or not 2 <= len(slides) <= 80:
        raise ValueError('slides must contain between 2 and 80 slide objects')
    normalized = []
    for index, raw in enumerate(slides, 1):
        if not isinstance(raw, dict):
            raise ValueError(f'slides[{index}] must be an object')
        title = _clean_text(raw.get('title'), field=f'slides[{index}].title', limit=500)
        bullets = _clean_text_list(raw.get('bullets'), field=f'slides[{index}].bullets', limit=20)
        equations = _clean_text_list(raw.get('equations'), field=f'slides[{index}].equations', limit=10)
        sources = _clean_text_list(raw.get('sources'), field=f'slides[{index}].sources', limit=30)
        notes = str(raw.get('notes') or '').strip()
        if len(notes) > 50000:
            raise ValueError(f'slides[{index}].notes exceeds 50000 characters')
        images = _clean_text_list(raw.get('images'), field=f'slides[{index}].images', limit=2)
        if not bullets and not equations and not images and index not in (1, len(slides)):
            raise ValueError(f'slides[{index}] needs bullets, equations, or images')
        if sum(len(item) for item in [*bullets, *equations]) > 1200:
            raise ValueError(f'slides[{index}] contains too much visible text; split it into more slides')
        normalized.append({
            'title': title, 'bullets': bullets, 'equations': equations,
            'sources': sources, 'notes': notes, 'images': images,
        })

    if template is not None and template.suffix.lower() != '.pptx':
        raise ValueError('scientific presentation templates must be .pptx files')
    prs = Presentation(str(template)) if template else Presentation()
    if not template:
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    else:
        # Preserve theme, masters, layouts, and page size but never reuse sample
        # slides as scientific evidence or mutate the uploaded template.
        for slide_id in list(prs.slides._sldIdLst):
            prs.part.drop_rel(slide_id.rId)
            prs.slides._sldIdLst.remove(slide_id)

    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
    layout = min(prs.slide_layouts, key=lambda item: abs(sum(
        ph.placeholder_format.type in body_types for ph in item.placeholders) - 1))
    fallback_body_count = 0
    inserted_images = 0
    missing_sources = []
    truncated_source_footers = []
    equations_as_text = 0
    resolved_images = image_paths or {}

    for page, spec in enumerate(normalized, 1):
        slide = prs.slides.add_slide(layout)
        title = slide.shapes.title
        if title is None:
            title = slide.shapes.add_textbox(
                int(prs.slide_width * .06), int(prs.slide_height * .04),
                int(prs.slide_width * .88), int(prs.slide_height * .13),
            )
        title.text = spec['title']

        has_images = bool(spec['images'])
        body = next((ph for ph in slide.placeholders
                     if ph.placeholder_format.type in body_types and ph.has_text_frame), None)
        if body is None:
            fallback_body_count += 1
            body = slide.shapes.add_textbox(
                int(prs.slide_width * .06), int(prs.slide_height * .21),
                int(prs.slide_width * (.52 if has_images else .88)), int(prs.slide_height * .64),
            )
        elif has_images:
            body.width = min(body.width, int(prs.slide_width * .52))
        frame = body.text_frame
        frame.clear()
        frame.word_wrap = True
        lines = [*spec['bullets'], *(f'公式：{item}' for item in spec['equations'])]
        equations_as_text += len(spec['equations'])
        for line_index, text in enumerate(lines):
            paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
            paragraph.text = text
            if not body.is_placeholder:
                paragraph.font.size = Pt(20)

        for image_index, name in enumerate(spec['images']):
            image = resolved_images.get(name)
            if image is None or not image.is_file():
                raise ValueError(f'Image attachment {name!r} is unavailable')
            top = int(prs.slide_height * (.23 + image_index * .31))
            picture = slide.shapes.add_picture(
                str(image), int(prs.slide_width * .62), top, width=int(prs.slide_width * .32),
            )
            max_height = int(prs.slide_height * .27)
            if picture.height > max_height:
                picture.width = int(picture.width * max_height / picture.height)
                picture.height = max_height
            inserted_images += 1

        if spec['sources']:
            footer = slide.shapes.add_textbox(
                int(prs.slide_width * .06), int(prs.slide_height * .90),
                int(prs.slide_width * .88), int(prs.slide_height * .055),
            )
            paragraph = footer.text_frame.paragraphs[0]
            source_text = '；'.join(spec['sources'])
            if len(source_text) > 300:
                source_text = source_text[:280].rstrip() + '…（完整来源见备注）'
                truncated_source_footers.append(page)
            paragraph.text = '来源：' + source_text
            paragraph.font.size = Pt(8)
            paragraph.alignment = PP_ALIGN.LEFT
        elif page not in (1, len(normalized)):
            missing_sources.append(page)

        note_sections = []
        if spec['notes']:
            note_sections.append(spec['notes'])
        if spec['sources']:
            note_sections.append('来源：\n' + '\n'.join(f'- {item}' for item in spec['sources']))
        if spec['equations']:
            note_sections.append(
                '公式说明：公式以可编辑文本写入，未转换为 PowerPoint 原生 OMML 对象。')
        if note_sections:
            slide.notes_slide.notes_text_frame.text = '\n\n'.join(note_sections)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    warnings = []
    if missing_sources:
        warnings.append({
            'code': 'MISSING_PAGE_SOURCES',
            'pages': missing_sources,
            'message': 'These content slides have no supplied source labels; verify them before use.',
        })
    if equations_as_text:
        warnings.append({
            'code': 'EQUATIONS_AS_TEXT',
            'count': equations_as_text,
            'message': 'Equations are editable text, not native PowerPoint OMML equations.',
        })
    if truncated_source_footers:
        warnings.append({
            'code': 'SOURCE_FOOTER_TRUNCATED',
            'pages': truncated_source_footers,
            'message': 'Visible source footers were shortened; complete source labels remain in speaker notes.',
        })
    return {
        'slide_count': len(normalized),
        'template_mode': 'layout' if template else 'none',
        'fallback_body_count': fallback_body_count,
        'inserted_images': inserted_images,
        'equations_as_text': equations_as_text,
        'warnings': warnings,
        'representation': 'native_pptx',
        'generation_status': 'completed',
        'visual_validation': 'not_performed',
    }
