import importlib.util
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Pt

spec = importlib.util.spec_from_file_location('pptx_template_fill', Path(__file__).parents[1] /
    'lazymind/chat/engine/tools/file_resources/pptx_template_fill.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


class TemplateFillTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.source = Path(self.temp.name) / 'template.pptx'
        self.output = Path(self.temp.name) / 'output.pptx'
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = 'Sample title'
        run = slide.placeholders[1].text_frame.paragraphs[0].add_run()
        run.text = 'Sample achievement'
        run.font.bold = True
        run.font.size = Pt(20)
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = 'Keep page'
        prs.save(self.source)
        self.pages = [module.inspect_template(self.source, p) for p in (1, 2)]
        self.edits = [{'id': s['id'], 'expected_text': s['text'], 'action': 'keep'}
                      for p in self.pages for s in p['slots']]

    def fill(self):
        return module.fill_template(self.source, self.output, self.pages[0]['sha256'], self.edits)

    def test_preserves_package_parts_and_styles(self):
        original = self.source.read_bytes()
        self.edits[0].update(action='replace', text='My <work> & results')
        result = self.fill()
        self.assertEqual(result['slide_count'], 2)
        self.assertEqual(self.source.read_bytes(), original)
        with ZipFile(self.source) as before, ZipFile(self.output) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'ppt/slides/slide1.xml':
                    self.assertEqual(before.read(name), after.read(name), name)
        actual = Presentation(self.output)
        self.assertEqual(actual.slides[0].shapes.title.text, 'My <work> & results')
        run = actual.slides[0].placeholders[1].text_frame.paragraphs[0].runs[0]
        self.assertTrue(run.font.bold)
        self.assertEqual(run.font.size, Pt(20))

    def test_missing_or_duplicate_slots_rejected(self):
        self.edits.pop()
        with self.assertRaises(ValueError):
            self.fill()
        self.assertFalse(self.output.exists())
        self.edits.append(self.edits[0])
        with self.assertRaises(ValueError):
            self.fill()

    def test_empty_placeholder_can_be_filled(self):
        empty = next(s for s in self.pages[1]['slots'] if s['empty_paragraph'])
        edit = next(e for e in self.edits if e['id'] == empty['id'])
        edit.update(action='replace', text='New work summary')
        self.fill()
        actual = Presentation(self.output)
        self.assertIn('New work summary', '\n'.join(s.text for s in actual.slides[1].shapes if s.has_text_frame))

    def test_stale_text_and_hash_rejected(self):
        self.edits[0]['expected_text'] = 'wrong'
        with self.assertRaises(ValueError):
            self.fill()
        with self.assertRaises(ValueError):
            module.fill_template(self.source, self.output, 'wrong', self.edits)

    def test_source_overwrite_rejected(self):
        with self.assertRaises(ValueError):
            module.fill_template(self.source, self.source, self.pages[0]['sha256'], self.edits)

    def test_overflow_warning_and_page_bounds(self):
        self.edits[0].update(action='replace', text='Long content ' * 30)
        self.assertTrue(self.fill()['warnings'])
        with self.assertRaises(ValueError):
            module.inspect_template(self.source, 0)


if __name__ == '__main__':
    unittest.main()
